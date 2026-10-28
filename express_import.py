"""Reviewed CSV/XLSX report imports from Express. No Express database writes."""
import json,io,csv,base64,hashlib,uuid,datetime,math,zipfile
from urllib.parse import parse_qs,urlsplit
import permissions_policy as policy
import budget_governance as bg
import finance_workspace as fw
TYPES={
 'Actuals':['date','kind','source','costCenter','amount','currency'],
 'Receivables':['reference','party','date','due','outstanding','currency'],
 'Payables':['reference','party','date','due','outstanding','currency'],
 'Chart of accounts':['account','name','category'],
 'General ledger':['date','reference','lineId','account','debit','credit','currency'],
 'Trial balance':['account','name','debit','credit','currency'],
 'Bank statement':['date','reference','account','debit','credit','currency'],
 'VAT register':['date','reference','party','taxId','amount','tax','currency'],
 'Withholding register':['date','reference','party','taxId','amount','tax','currency'],
 'Fixed assets':['account','name','date','cost','depreciation','netBookValue','currency']}
ALIASES={'date':['วันที่','วันที่เอกสาร'],'reference':['เลขที่เอกสาร','เลขที่','document','docno'],'party':['ลูกค้า','เจ้าหนี้','ชื่อลูกค้า','ชื่อผู้จำหน่าย'],'account':['รหัสบัญชี','เลขที่บัญชี'],'name':['ชื่อบัญชี','ชื่อ'],'debit':['เดบิต'],'credit':['เครดิต'],'amount':['จำนวนเงิน','ยอดเงิน'],'outstanding':['ยอดคงเหลือ','ค้างชำระ','balance'],'due':['ครบกำหนด','วันที่ครบกำหนด'],'tax':['ภาษี','vat'],'taxId':['เลขประจำตัวผู้เสียภาษี'],'costCenter':['แผนก','ศูนย์ต้นทุน'],'source':['แบรนด์','แหล่งรายได้'],'currency':['สกุลเงิน']}
def initialize(c):
 c.execute('CREATE TABLE IF NOT EXISTS express_batches(id TEXT PRIMARY KEY,payload TEXT NOT NULL,version INTEGER NOT NULL)')
 c.execute('CREATE TABLE IF NOT EXISTS express_profiles(name TEXT PRIMARY KEY,payload TEXT NOT NULL)')
def decode_file(body):
 filename=str(body.get('filename') or '');raw=base64.b64decode(body.get('file',''),validate=True)
 if not 0<len(raw)<=10000000:raise ValueError('Choose a CSV/XLSX file up to 10 MB')
 header=int(body.get('headerRow',1))
 if not 1<=header<=100:raise ValueError('Header row must be 1–100')
 if filename.lower().endswith('.csv'):
  encoding=body.get('encoding') or 'utf-8-sig'
  if encoding not in ['utf-8-sig','cp874','utf-16']:raise ValueError('Choose UTF-8, Thai Windows-874 or UTF-16')
  text=raw.decode(encoding);delimiter=body.get('delimiter') or ','
  if delimiter not in [',',';','\t']:raise ValueError('Invalid delimiter')
  import itertools
  matrix=list(itertools.islice(csv.reader(io.StringIO(text),delimiter=delimiter),11001))
 elif filename.lower().endswith('.xlsx'):
  with zipfile.ZipFile(io.BytesIO(raw)) as z:
   if sum(x.file_size for x in z.infolist())>50000000:raise ValueError('Workbook expands beyond the safe size limit')
   import re
   if any(re.search(rb'<(?:\w+:)?f(?:\s|>)',z.read(x)) for x in z.namelist() if x.startswith('xl/worksheets/') and x.endswith('.xml')):raise ValueError('Save formula results as values before importing XLSX, or use the Express CSV export')
  import openpyxl
  book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
  sheet=body.get('sheet') or book.sheetnames[0]
  if sheet not in book.sheetnames:raise ValueError('Sheet not found')
  ws=book[sheet]
  if ws.max_row>11000 or ws.max_column>150:raise ValueError('Limit each report to 10,000 rows and 150 columns')
  matrix=[list(r) for r in ws.iter_rows(values_only=True)];book.close()
 else:raise ValueError('Export from Express as CSV or save the report as XLSX; legacy XLS and backup/database files are not accepted')
 if len(matrix)<header or len(matrix)>11000:raise ValueError('Choose the correct header row; limit to 10,000 data rows')
 headings=[str(x or '').strip() for x in matrix[header-1]]
 if not headings or len(headings)>150 or any(not x for x in headings) or len(set(headings))!=len(headings):raise ValueError('The selected header row needs unique, nonblank column names; remove merged/blank report columns first')
 rows=[dict(zip(headings,[str(v.isoformat() if isinstance(v,(datetime.date,datetime.datetime)) else v if v is not None else '') for v in row])) for row in matrix[header:] if any(v is not None and str(v).strip() for v in row)]
 if len(rows)>10000:raise ValueError('Limit report to 10,000 rows')
 return headings,rows,hashlib.sha256(raw).hexdigest()
def date(value,calendar="Auto"):
 s=str(value).strip()[:10]
 import re
 short=re.fullmatch(r'(\d{1,2})[/-](\d{1,2})[/-](\d{2})',s)
 if short:
  if calendar not in ['CE','BE']:raise ValueError('Choose CE or Thai BE for two-digit-year dates')
  dd,mm,yy=map(int,short.groups());s=f'{dd:02d}/{mm:02d}/{(2500 if calendar=="BE" else 2000)+yy}'
 try:d=datetime.date.fromisoformat(s)
 except ValueError:
  d=None
  for fmt in ['%d/%m/%Y','%d-%m-%Y']:
   try:d=datetime.datetime.strptime(s,fmt).date();break
   except ValueError:pass
  if d is None:raise ValueError('Use YYYY-MM-DD or DD/MM/YYYY dates')
 if d.year>2400:d=d.replace(year=d.year-543)
 if not 1900<=d.year<=2200:raise ValueError('Date year is outside supported range')
 return d.isoformat()
def numeric(v):
 s=str(v).strip().replace(',','');s='-'+s[1:-1] if s.startswith('(') and s.endswith(')') else s
 n=float(s or 0)
 if not math.isfinite(n) or abs(n)>1e15:raise ValueError('Invalid financial amount')
 return n
def normalized(kind,rows,mapping,defaults,year,calendar="Auto"):
 clean=[];errors=[];seen=set()
 for index,row in enumerate(rows,1):
  try:
   r={field:str(row.get(mapping.get(field),defaults.get(field,''))).strip() for field in TYPES[kind]}
   for field in TYPES[kind]:
    if field in ['date','due']:r[field]=date(r[field],calendar)
    elif field in ['amount','outstanding','tax','debit','credit','cost','depreciation','netBookValue']:
     if not r[field] and field not in ['debit','credit']:raise ValueError('Missing '+field)
     r[field]=numeric(r[field])
    elif not r[field] and field not in ['category','taxId','name']:raise ValueError('Missing '+field)
   if 'currency' in r and r['currency'] not in ['THB','USD','EUR','CNY']:raise ValueError('Unmapped currency')
   if kind=='Actuals':
    if r['kind'] not in ['Revenue','COGS','Expense'] or r['currency']!='THB':raise ValueError('Map kind to Revenue/COGS/Expense and convert to THB before importing actuals')
    d=datetime.date.fromisoformat(r['date']);fy=d.year+int(d.month>=5)
    if fy!=int(year):raise ValueError('Actual date falls outside the selected May–April fiscal year')
   for field in ['outstanding','debit','credit','cost','depreciation','netBookValue']:
    if r.get(field,0)<0:raise ValueError(field+' must be nonnegative')
   if kind=='Fixed assets' and abs(r['cost']-r['depreciation']-r['netBookValue'])>.01:raise ValueError('Asset cost minus accumulated depreciation must equal net book value')
   key=json.dumps(r,sort_keys=True)
   if key in seen:raise ValueError('Exact duplicate data row')
   seen.add(key);clean.append(r)
  except (ValueError,TypeError) as ex:errors.append({'row':index,'error':str(ex)})
 if kind in ['General ledger','Trial balance'] and not errors:
  groups={}
  for r in clean:
   key=(r['currency'],r.get('date'),r.get('reference')) if kind=='General ledger' else (r['currency'],)
   groups[key]=groups.get(key,0)+r['debit']-r['credit']
  if any(abs(v)>.01 for v in groups.values()):errors.append({'row':0,'error':'Debits and credits do not balance by currency/document. Export the complete journal or trial balance.'})
 return clean,errors

def apply(c,batch,actor):
 kind=batch['type'];rows=batch['rows'];effects={}
 if kind=='Actuals':
  data,state=bg.workspace(c,batch['year']);before=json.loads(json.dumps(data));lines=bg.grouped(state);totals={}
  for r in rows:
   key=(r['kind'],r['costCenter'],r['source']);m=(datetime.date.fromisoformat(r['date']).month-5)%12;totals[(key,m)]=totals.get((key,m),0)+r['amount']
  effects['actuals']=[]
  for (key,m),total in totals.items():
   line=next((l for l in lines if (l['kind'],l['costCenter'],l['source'])==key),None)
   if not line:raise ValueError('Map '+str(key)+' to an existing approved Financial source before applying')
   effects['actuals'].append({'key':list(key),'month':m,'before':line['actual'][m],'after':total});line['actual'][m]=total
  for line in lines:line['target']=line['actual']
  bg.publish(c,batch['year'],lines,actual=True);after,_=bg.workspace(c,batch['year']);bg.guard(c,before,after)
 elif kind in ['Receivables','Payables']:
  f,version=fw.load(c,batch['year']);sources,_=fw.source_rows(c,json.loads(json.dumps(f)));effects['manualIds']=[];effects['updatedManual']={};effects['overrides']={};f.setdefault('expressBalances',{})
  typ='Receivable' if kind=='Receivables' else 'Payable'
  for r in rows:
   candidates=[x for x in sources if x['reference'].casefold()==r['reference'].casefold() and x['currency']==r['currency'] and x['type']==typ]
   if len(candidates)>1:raise ValueError('Ambiguous reference '+r['reference'])
   if candidates:
    source=candidates[0]
    import unicodedata
    norm=lambda v:' '.join(unicodedata.normalize('NFKC',str(v)).casefold().split())
    if norm(source['party'])!=norm(r['party']):raise ValueError('Customer/supplier name differs for '+r['reference']+'; reconcile the master-name mapping before importing')
    if not source['id'].startswith('IS:'):
     original=next(x for x in f['manual'] if x['id']==source['id'])
     if not original.get('expressBatch'):raise ValueError('Reference already exists as a manual Financial balance: '+r['reference'])
     if original.get('asOf','')>=batch['asOf']:raise ValueError('A same-date or newer Express balance exists; reverse corrections in order')
     before=json.loads(json.dumps(original));original.update(amount=r['outstanding'],asOf=batch['asOf'],due=r['due'],expressBatch=batch['id'],origin='Express outstanding snapshot at '+batch['asOf'])
     effects['updatedManual'][original['id']]={'before':before,'after':json.loads(json.dumps(original))};continue
    if f['expressBalances'].get(source['id'],{}).get('asOf','')>=batch['asOf']:raise ValueError('A same-date or newer IS/Express reconciliation exists; reverse the newer report first')
    if r['outstanding']>source['amount']+.01:raise ValueError('Express outstanding exceeds the matching IS invoice total: '+r['reference'])
    effects['overrides'][source['id']]={'before':f['expressBalances'].get(source['id']),'after':{'outstanding':r['outstanding'],'asOf':batch['asOf'],'batch':batch['id']}}
    f['expressBalances'][source['id']]=effects['overrides'][source['id']]['after']
   else:
    ident='EXP:'+str(uuid.uuid4());f['manual'].append({'id':ident,'type':typ,'reference':r['reference'],'party':r['party'],'date':r['date'],'due':r['due'],'currency':r['currency'],'amount':r['outstanding'],'asOf':batch['asOf'],'origin':'Express outstanding snapshot at '+batch['asOf'],'expressBatch':batch['id'],'at':bg.now(),'by':actor[4]});effects['manualIds'].append(ident)
  f['history'].append({'id':str(uuid.uuid4()),'at':bg.now(),'by':actor[4],'name':actor[2],'action':'Express import approved','details':{'batch':batch['id'],'type':kind}});c.execute('insert or replace into finance_workspace values(?,?,?)',('company',json.dumps(f,ensure_ascii=False),version+1))
 return effects

def reverse(c,batch,actor):
 effects=batch.get('effects',{})
 if effects.get('actuals'):
  data,state=bg.workspace(c,batch['year']);before=json.loads(json.dumps(data));lines=bg.grouped(state)
  for item in effects['actuals']:
   line=next(l for l in lines if [l['kind'],l['costCenter'],l['source']]==item['key']);m=item['month']
   if abs(line['actual'][m]-item['after'])>.01:raise ValueError('Imported actuals changed later; reconcile them before reversal')
   line['actual'][m]=item['before']
  for l in lines:l['target']=l['actual']
  bg.publish(c,batch['year'],lines,actual=True);after,_=bg.workspace(c,batch['year']);bg.guard(c,before,after)
 if 'manualIds' in effects:
  f,version=fw.load(c,batch['year'])
  if any(x['sourceId'] in effects['manualIds'] and not x.get('void') for x in f['settlements']):raise ValueError('Reverse related settlements before reversing the imported obligation')
  for ident,change in effects.get('updatedManual',{}).items():
   current=next(r for r in f['manual'] if r['id']==ident)
   if current!=change['after']:raise ValueError('A later balance update exists; reverse the newer import first')
   current.clear();current.update(change['before'])
  for r in f['manual']:
   if r['id'] in effects['manualIds'] and r.get('expressBatch')!=batch['id']:raise ValueError('A later Express snapshot updated this obligation; reverse it first')
  for ident,change in effects['overrides'].items():
   if f.get('expressBalances',{}).get(ident)!=change['after']:raise ValueError('A later Express snapshot superseded this import; reverse the newer import first')
   if change['before'] is None:f['expressBalances'].pop(ident,None)
   else:f['expressBalances'][ident]=change['before']
  for r in f['manual']:
   if r['id'] in effects['manualIds']:r['void']={'at':bg.now(),'by':actor[4],'reason':'Express batch reversed'}
  c.execute('insert or replace into finance_workspace values(?,?,?)',('company',json.dumps(f,ensure_ascii=False),version+1))

def handle(h,method,connect,session,permissions,origin):
 actor=session(h)
 if not actor:return h.reply(401,{'error':'Sign in required'})
 try:
  perms=permissions(actor)
  for key in ['budget_view','budget.finance.view','budget.express.view']:policy.need(perms,key)
  year=parse_qs(urlsplit(h.path).query).get('year',[''])[0]
  if not year.isdigit() or not 2000<=int(year)<=2200:raise ValueError('Choose a fiscal year')
  with connect() as c:
   c.execute('BEGIN IMMEDIATE');initialize(c)
   selected=parse_qs(urlsplit(h.path).query).get('batch',[None])[0]
   if method=='GET' and selected:
    row=c.execute('select payload from express_batches where id=?',(selected,)).fetchone()
    if not row:raise ValueError('Import batch not found')
    batch=json.loads(row[0]);return h.reply(200,{'type':batch['type'],'asOf':batch['asOf'],'rows':batch['rows']})
   if method=='POST':
    if not origin(h):return h.reply(403,{'error':'Invalid origin'})
    length=int(h.headers.get('Content-Length',0))
    if not 0<length<=15000000:raise ValueError('Invalid upload size')
    body=json.loads(h.rfile.read(length));action=body.get('action');policy.need(perms,'budget.finance.review' if action in ['approve','reverse','reject'] else 'budget.finance.import')
    if action in ['inspect','stage']:
     kind=body.get('type')
     if kind not in TYPES:raise ValueError('Choose a report type')
     headers,raw,digest=decode_file(body)
     if action=='inspect':return h.reply(200,{'headers':headers,'sample':raw[:10],'count':len(raw),'fields':TYPES[kind],'suggested':{f:next((x for x in headers if x.casefold() in [f.casefold()]+[a.casefold() for a in ALIASES.get(f,[])]),'') for f in TYPES[kind]}})
     company=str(body.get('company') or '').strip();asof=date(body.get('asOf'))
     if datetime.date.fromisoformat(asof)>fw.today():raise ValueError('Choose an as-of date today or earlier')
     if not company or len(company)>200:raise ValueError('Enter the legal company name/code used in this Express export')
     existing=list(c.execute('select payload from express_batches'));active=[json.loads(r[0]) for r in existing if json.loads(r[0])['status'] not in ['Rejected','Reversed']]
     if any(r['company']!=company for r in active):raise ValueError('This installation has a different legal-company import scope; use a separate workspace for another legal entity')
     if any(r['digest']==digest and r['type']==kind for r in active):raise ValueError('This file/report is already staged or imported')
     rows,errors=normalized(kind,raw,body.get('mapping') or {},body.get('defaults') or {},year,body.get('dateCalendar','Auto'))
     if errors:return h.reply(400,{'error':'Resolve validation errors before staging','errors':errors[:100]})
     if not rows:raise ValueError('No report data rows')
     ident=str(uuid.uuid4());batch={'id':ident,'type':kind,'year':year,'company':company,'asOf':asof,'filename':str(body.get('filename'))[:200],'digest':digest,'status':'Staged','submittedBy':str(actor[4]),'submittedName':actor[2],'at':bg.now(),'rows':rows,'history':[]}
     c.execute('insert into express_batches values(?,?,1)',(ident,json.dumps(batch,ensure_ascii=False)))
     profile=str(body.get('profile') or '').strip()
     if profile:c.execute('insert or replace into express_profiles values(?,?)',(profile[:100],json.dumps({k:body.get(k) for k in ['type','mapping','defaults','company','encoding','delimiter','headerRow','sheet','dateCalendar']})))
    elif action in ['approve','reject','reverse']:
     record=c.execute('select payload,version from express_batches where id=?',(body.get('id'),)).fetchone()
     if not record or record[1]!=body.get('version'):return h.reply(409,{'error':'Import changed; refresh before reviewing'})
     batch=json.loads(record[0]);reason=str(body.get('reason') or '').strip()
     if action in ['approve','reject']:
      if batch['status']!='Staged':raise ValueError('This batch is not awaiting review')
      if batch['submittedBy']==str(actor[4]):raise policy.PermissionDenied('Another authorized reviewer must decide your import')
     if action=='approve':
      for key in ['is.invoice.view','is.purchase-order.view']:policy.need(perms,key)
      for payload, in c.execute('select payload from express_batches where id!=?',(batch['id'],)):
       previous=json.loads(payload)
       if previous['status']=='Imported' and all(previous[k]==batch[k] for k in ['company','type','asOf']):raise ValueError('An active report of this type and as-of date already exists; reverse it before importing a correction')
      if batch['type']=='Actuals':
       for r in batch['rows']:policy.need(perms,'budget.'+('products' if r['kind']=='Revenue' else 'cogs' if r['kind']=='COGS' else policy.CATEGORY.get(r['source'],'manage_items'))+('.edit' if r['kind']=='Expense' and r['source'] not in policy.CATEGORY else '.actual'))
      batch['effects']=apply(c,batch,actor);batch['status']='Imported'
     elif action=='reject':
      if not reason:raise ValueError('A rejection reason is required')
      batch['status']='Rejected'
     else:
      if batch['status']!='Imported' or not reason:raise ValueError('Only imported batches may be reversed; a reason is required')
      if batch['type']=='Actuals':
       for r in batch['rows']:policy.need(perms,'budget.'+('products' if r['kind']=='Revenue' else 'cogs' if r['kind']=='COGS' else policy.CATEGORY.get(r['source'],'manage_items'))+('.edit' if r['kind']=='Expense' and r['source'] not in policy.CATEGORY else '.actual'))
      reverse(c,batch,actor);batch['status']='Reversed'
     batch['history'].append({'at':bg.now(),'by':actor[4],'name':actor[2],'action':action,'reason':reason});c.execute('update express_batches set payload=?,version=version+1 where id=?',(json.dumps(batch,ensure_ascii=False),batch['id']))
    else:raise ValueError('Unknown Express action')
   batches=[]
   for payload,version in c.execute('select payload,version from express_batches order by rowid desc limit 100'):
    b=json.loads(payload);b['version']=version;b['rowCount']=len(b['rows']);b['totalsByCurrency']={currency:{field:round(sum(float(r.get(field,0)) for r in b['rows'] if r.get('currency','Codes')==currency),2) for field in ['amount','outstanding','debit','credit','tax','cost','depreciation','netBookValue']} for currency in {r.get('currency','Codes') for r in b['rows']}};b['rows']=b['rows'][:500];batches.append(b)
   return h.reply(200,{'batches':batches,'profiles':{name:json.loads(p) for name,p in c.execute('select name,payload from express_profiles')},'types':TYPES,'permissions':sorted(perms),'userId':str(actor[4])})
 except policy.PermissionDenied as ex:return h.reply(403,{'error':str(ex)})
 except (ValueError,TypeError,KeyError,StopIteration,UnicodeError,zipfile.BadZipFile) as ex:return h.reply(400,{'error':str(ex) or 'Invalid report'})
